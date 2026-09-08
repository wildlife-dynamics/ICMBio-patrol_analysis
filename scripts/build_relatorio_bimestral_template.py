"""Convert the partner's real Relatório Bimestral mockup into a docxtpl
(Jinja2) template, following the same tag style as APN's real, working
patrol/event report template (ecoscope-workflows-ext-apn's
generate_patrol_report / _report.py).

Run once (or whenever the partner supplies a revised mockup):
    pixi run -e default python scripts/build_relatorio_bimestral_template.py

What this does, section by section (matching the mockup's own layout,
read from /Users/zak/Downloads/Relatório Bimestral_modelo.docx):
  - Appends `{{ var }}` after each stat line (Período, Número de patrulhas,
    Quilômetros percorridos, Participantes/operadores).
  - Replaces the 3 static map screenshots already embedded in the mockup
    with `{{ patrol_map }}` / `{{ time_density_map }}` / `{{ events_map }}`
    (docxtpl InlineImage placeholders), in the order they appear.
  - Turns each existing table's single blank data row into a docxtpl
    row-loop (`{%tr for item in ... %}` / data row / `{%tr endfor %}`),
    the same 3-row pattern APN's template uses — header row and column
    labels are left exactly as the partner wrote them.
  - Appends two NEW sections at the end (Atropelamentos, Resumo por setor)
    not present in the partner's mockup — see spec.yaml's "Roadkill Events
    Table" / "Grid Summary Table" comments for why these exist: roadkill
    events don't fit the Ameaças columns, and the PRD explicitly asks for
    a grid-square summary table that the mockup omits.

Column-name notes (kept in code, not silently baked into the partner's
docx):
  - Camera traps table: the mockup's header literally says "Local", but
    our real data/column is called "Setor" (matches the PRD wording and
    spec.yaml's Sector Layer). The header text is left as "Local" here;
    ecoscope_workflows_ext_icmbio.tasks.prepare_bimonthly_report_context
    maps our `Setor` column to the `Local` key expected by this template.
  - The mockup's camera-traps table has no "Tipo de evento" column at all
    (PRD text lists 5 columns; the mockup only has 4: Local, Data,
    Responsável, Coordenadas) — not added here since this script only
    converts what's literally in the mockup; flag to the partner if the
    5th column is actually wanted.
  - The "Ameaças" table has a 4th, entirely empty header/data column in
    the mockup (both header and sample row are blank there) — left empty
    in every row rather than guessing at its purpose or deleting a column
    (safer than column-count XML surgery).
"""

import copy
from pathlib import Path

import docx
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

SRC = Path("/Users/zak/Downloads/Relatório Bimestral_modelo.docx")
DST = Path(__file__).parent.parent / "resources" / "templates" / "relatorio_bimestral_template.docx"


# ── low-level helpers ─────────────────────────────────────────────────────────

def find_paragraph(doc: docx.Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(f"paragraph not found (prefix={prefix!r})")


def append_jinja(paragraph: Paragraph, expr: str) -> None:
    ref = paragraph.runs[0] if paragraph.runs else None
    run = paragraph.add_run(expr)
    if ref is not None:
        run.bold = ref.bold
        run.font.size = ref.font.size
        run.font.name = ref.font.name


def clear_cell(cell) -> Paragraph:
    cell.text = ""
    return cell.paragraphs[0]


def set_cell_jinja(cell, expr: str, bold: bool = False):
    p = clear_cell(cell)
    r = p.add_run(expr)
    r.bold = bold
    return r


def image_paragraphs(doc: docx.Document) -> list[Paragraph]:
    """Body paragraphs containing an embedded drawing, in document order."""
    out = []
    for p in doc.paragraphs:
        if p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}graphic"):
            out.append(p)
    return out


def replace_image_with_jinja(paragraph: Paragraph, expr: str) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(expr)


def clone_row(table: Table, ref_row, after: bool) -> _Row:
    new_tr = copy.deepcopy(ref_row._tr)
    if after:
        ref_row._tr.addnext(new_tr)
    else:
        ref_row._tr.addprevious(new_tr)
    return _Row(new_tr, table)


def make_row_loop(table: Table, data_row_index: int, loop_var: str, cell_exprs: list[str]) -> None:
    """Turn table.rows[data_row_index] (a single blank data row) into a
    docxtpl row-loop: a loop-start row, the data row (filled with
    cell_exprs), and a loop-end row — the same 3-row pattern used by
    APN's real, working template (ecoscope_workflows_ext_apn's Eventos /
    IMAGENES tables)."""
    data_row = table.rows[data_row_index]
    ncols = len(data_row.cells)

    loop_start = clone_row(table, data_row, after=False)
    loop_end = clone_row(table, data_row, after=True)

    set_cell_jinja(loop_start.cells[0], "{%tr for item in " + loop_var + " %}")
    for c in loop_start.cells[1:]:
        clear_cell(c)

    for cell, expr in zip(data_row.cells, cell_exprs):
        set_cell_jinja(cell, expr)
    for cell in data_row.cells[len(cell_exprs):]:
        clear_cell(cell)

    set_cell_jinja(loop_end.cells[0], "{%tr endfor %}")
    for c in loop_end.cells[1:]:
        clear_cell(c)

    assert len(data_row.cells) == ncols


def add_section_heading(doc: docx.Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.underline = True
    from docx.shared import Pt
    r.font.size = Pt(12)


# ── main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    doc = docx.Document(str(SRC))

    # ── stat lines ────────────────────────────────────────────────────────────
    append_jinja(find_paragraph(doc, "Período:"), "{{ period }}")
    append_jinja(find_paragraph(doc, "Número de patrulhas:"), " {{ total_patrols }}")
    append_jinja(find_paragraph(doc, "Quilômetros percorridos:"), "{{ total_distance_km }}")
    append_jinja(find_paragraph(doc, "Participantes/operadores:"), " {{ total_participants }}")

    # ── 3 maps, in document order: patrol map, time density, events map ──────
    imgs = image_paragraphs(doc)
    if len(imgs) != 3:
        raise RuntimeError(f"expected exactly 3 image paragraphs in the mockup, found {len(imgs)}")
    replace_image_with_jinja(imgs[0], "{{ patrol_map }}")
    replace_image_with_jinja(imgs[1], "{{ time_density_map }}")
    replace_image_with_jinja(imgs[2], "{{ events_map }}")

    # ── 5 existing tables, in document order ──────────────────────────────────
    tables = doc.tables
    if len(tables) != 5:
        raise RuntimeError(f"expected exactly 5 tables in the mockup, found {len(tables)}")

    patrols_tbl, events_summary_tbl, camera_traps_tbl, threats_tbl, photos_tbl = tables

    make_row_loop(
        patrols_tbl, 1, "patrols",
        [
            "{{ item['ID da patrulha'] }}",
            "{{ item['Início'] }}",
            "{{ item['Fim'] }}",
            "{{ item['Distância percorrida (km)'] }}",
        ],
    )

    make_row_loop(
        events_summary_tbl, 1, "events_summary",
        ["{{ item['Tipo de evento'] }}", "{{ item['Quantidade'] }}"],
    )

    # Mockup's header literally reads "Local" (not "Setor") — kept as-is;
    # ecoscope_workflows_ext_icmbio.tasks.prepare_bimonthly_report_context
    # maps our real `Setor` column onto this `Local` key.
    make_row_loop(
        camera_traps_tbl, 1, "camera_traps",
        [
            "{{ item['Local'] }}",
            "{{ item['Data'] }}",
            "{{ item['Responsável'] }}",
            "{{ item['Coordenadas'] }}",
        ],
    )

    # 4th column of this table is blank in the mockup itself (header and
    # sample row) — left blank in every generated row too.
    make_row_loop(
        threats_tbl, 1, "threats",
        [
            "{{ item['ID do evento'] }}",
            "{{ item['Tipo de evento'] }}",
            "{{ item['Detalhes do evento'] }}",
        ],
    )

    make_row_loop(
        photos_tbl, 1, "photos",
        ["{{ item.info }}", "{{ item.image }}"],
    )

    # ── new sections, not in the partner's mockup ─────────────────────────────
    # See spec.yaml's "Roadkill Events Table" / "Grid Summary Table" comments:
    # atropelamento_v2 events carry none of the Ameaças detail fields (own
    # table needed), and the PRD explicitly asks for a grid-square summary
    # the mockup doesn't include.
    add_section_heading(doc, "Atropelamentos")
    roadkill_tbl = doc.add_table(rows=2, cols=3)
    for i, h in enumerate(["ID do evento", "Tipo de evento", "Detalhes do evento"]):
        set_cell_jinja(roadkill_tbl.rows[0].cells[i], h, bold=True)
    make_row_loop(
        roadkill_tbl, 1, "roadkill",
        [
            "{{ item['ID do evento'] }}",
            "{{ item['Tipo de evento'] }}",
            "{{ item['Detalhes do evento'] }}",
        ],
    )

    add_section_heading(doc, "Resumo por setor")
    grid_tbl = doc.add_table(rows=2, cols=3)
    for i, h in enumerate(["Grupo", "Quantidade de eventos", "Coordenadas centrais"]):
        set_cell_jinja(grid_tbl.rows[0].cells[i], h, bold=True)
    make_row_loop(
        grid_tbl, 1, "grid_summary",
        [
            "{{ item['Grupo'] }}",
            "{{ item['Quantidade de eventos'] }}",
            "{{ item['Coordenadas centrais'] }}",
        ],
    )

    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DST))
    print(f"Template written to: {DST}")


if __name__ == "__main__":
    main()
